"""Reporting periods, the PDF and Excel renderers and the export endpoint's guards.

The reports here are fixtures built in the test; the database-backed exports
are in ``test_reports_database.py``.
"""

import datetime as dt
import io
import re
from decimal import Decimal
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from principals import TESTER_NAME, as_user
from pydantic import ValidationError

from app.auth import sessions
from app.core.authorization import ANONYMOUS, get_user_principal
from app.core.config import get_settings
from app.core.permissions import Permission as P
from app.main import create_app
from app.reports import router as report_router
from app.reports.catalog import REPORTS
from app.reports.excel import render_excel
from app.reports.model import Chart, Column, Kpi, Report, Series, Table
from app.reports.pdf import render_pdf
from app.reports.period import PeriodError, PeriodIn, resolve

TODAY = dt.date(2026, 10, 9)
D = dt.date


# --- Periods --------------------------------------------------------------------


@pytest.mark.parametrize(
    ("mode", "today", "start", "end", "label"),
    [
        ("this-month", TODAY, D(2026, 10, 1), D(2026, 10, 31), "This Month · October 2026"),
        ("last-month", TODAY, D(2026, 9, 1), D(2026, 9, 30), "Last Month · September 2026"),
        ("ytd", TODAY, D(2026, 1, 1), TODAY, "Year to Date · Jan 1 – Oct 9, 2026"),
        (
            "last-3-months",
            TODAY,
            D(2026, 8, 1),
            D(2026, 10, 31),
            "Last 3 Months · Aug 1 – Oct 31, 2026",
        ),
        (
            "last-6-months",
            TODAY,
            D(2026, 5, 1),
            D(2026, 10, 31),
            "Last 6 Months · May 1 – Oct 31, 2026",
        ),
        ("this-year", TODAY, D(2026, 1, 1), D(2026, 12, 31), "This Year · 2026"),
        # January: last month and rolling windows reach into the previous year.
        (
            "last-month",
            D(2026, 1, 15),
            D(2025, 12, 1),
            D(2025, 12, 31),
            "Last Month · December 2025",
        ),
        (
            "last-3-months",
            D(2026, 1, 15),
            D(2025, 11, 1),
            D(2026, 1, 31),
            "Last 3 Months · Nov 1, 2025 – Jan 31, 2026",
        ),
        ("ytd", D(2026, 1, 1), D(2026, 1, 1), D(2026, 1, 1), "Year to Date · Jan 1, 2026"),
        # December: the month and the year end together.
        (
            "this-month",
            D(2026, 12, 31),
            D(2026, 12, 1),
            D(2026, 12, 31),
            "This Month · December 2026",
        ),
        (
            "ytd",
            D(2026, 12, 31),
            D(2026, 1, 1),
            D(2026, 12, 31),
            "Year to Date · Jan 1 – Dec 31, 2026",
        ),
        (
            "last-month",
            D(2026, 12, 31),
            D(2026, 11, 1),
            D(2026, 11, 30),
            "Last Month · November 2026",
        ),
        # Leap year February.
        ("this-month", D(2028, 2, 10), D(2028, 2, 1), D(2028, 2, 29), "This Month · February 2028"),
    ],
)
def test_presets(mode: str, today: dt.date, start: dt.date, end: dt.date, label: str) -> None:
    period = resolve(mode, today)
    assert (period.start, period.end) == (start, end)
    assert period.full_label == label


def test_specific_year_month_and_custom_range() -> None:
    year = resolve("year", TODAY, year=2025)
    assert (year.start, year.end, year.full_label) == (D(2025, 1, 1), D(2025, 12, 31), "2025")
    month = resolve("month", TODAY, year=2026, month=8)
    assert (month.start, month.end) == (D(2026, 8, 1), D(2026, 8, 31))
    assert month.full_label == "August 2026 · Aug 1 – Aug 31, 2026"
    december = resolve("month", TODAY, year=2025, month=12)
    assert december.end == D(2025, 12, 31)
    custom = resolve("custom", TODAY, start=D(2025, 11, 15), end=D(2026, 2, 3))
    assert custom.full_label == "Custom Range · Nov 15, 2025 – Feb 3, 2026"
    assert custom.year is None
    everything = resolve("all", TODAY)
    assert (everything.start, everything.end, everything.full_label) == (None, None, "All dates")


@pytest.mark.parametrize(
    ("mode", "kwargs", "part"),
    [
        ("month", {"year": 2026, "month": 8}, "2026-08"),
        ("ytd", {}, "2026-YTD"),
        ("year", {"year": 2025}, "2025"),
        ("this-month", {}, "2026-10"),
        ("custom", {"start": D(2026, 1, 1), "end": D(2026, 3, 31)}, "2026-01-01_2026-03-31"),
        ("all", {}, "All-Dates"),
    ],
)
def test_filename_parts(mode: str, kwargs: dict[str, Any], part: str) -> None:
    assert resolve(mode, TODAY, **kwargs).filename_part == part


@pytest.mark.parametrize(
    ("mode", "kwargs"),
    [
        ("year", {}),
        ("month", {"year": 2026}),
        ("month", {"year": 2026, "month": 13}),
        ("custom", {"start": D(2026, 3, 1)}),
        ("custom", {"start": D(2026, 3, 1), "end": D(2026, 2, 1)}),
        ("year", {"year": 1999}),
        ("someday", {}),
    ],
)
def test_invalid_periods(mode: str, kwargs: dict[str, Any]) -> None:
    with pytest.raises(PeriodError):
        resolve(mode, TODAY, **kwargs)


def test_period_input_is_validated() -> None:
    assert PeriodIn.model_validate({"mode": "month", "year": 2026, "month": 1}).resolve(
        TODAY
    ).start == D(2026, 1, 1)
    for bad in (
        {"mode": "month", "year": 2026},
        {"mode": "custom", "from": "2026-02-01", "to": "2026-01-01"},
        {"mode": "ytd", "extra": 1},
        {"mode": "decade"},
    ):
        with pytest.raises(ValidationError):
            PeriodIn.model_validate(bad)


# --- Renderers --------------------------------------------------------------------


def _report(rows: int = 3) -> Report:
    return Report(
        key="fixture",
        title="Fixture Report",
        module="Quality · Fixtures",
        period_label="Year to Date · Jan 1 – Oct 9, 2026",
        filename="Operations-Platform_Fixture_2026-YTD.pdf",
        filters=[("Department", "Production"), ("Source", "Internal Audit")],
        generated_at=dt.datetime(2026, 10, 9, 22, 30, tzinfo=dt.UTC),
        generated_by=TESTER_NAME,
        kpis=[
            Kpi("Open", 12),
            Kpi("Cost", Decimal("1234.5"), "money"),
            Kpi("Share", Decimal("0.125"), "percent"),
            Kpi("Not recorded", None, "money"),
        ],
        charts=[
            Chart("By month", ["Jan", "Feb", "Mar"], [Series("Opened", [1, None, 3])]),
            Chart(
                "By department",
                [f"Department {i}" for i in range(20)],
                [Series("CARs", list(range(20)))],
                kind="hbar",
            ),
            Chart(
                "Rate",
                ["Jan", "Feb"],
                [
                    Series("YTD", [Decimal("1.5"), Decimal("2")]),
                    Series("12MRA", [None, Decimal("1.1")]),
                ],
                kind="line",
            ),
            Chart("Empty", [], []),
        ],
        tables=[
            Table(
                "Register",
                [
                    Column("Number", "text"),
                    Column("Date", "date"),
                    Column("Cost", "money"),
                    Column("Share", "percent"),
                    Column("Days", "int"),
                ],
                [
                    [
                        f"CAR-{i:04d}",
                        D(2026, 1, 1) + dt.timedelta(days=i),
                        Decimal("10.25") if i % 2 else None,
                        Decimal("0.5"),
                        i,
                    ]
                    for i in range(rows)
                ],
                sheet="Register",
                pdf_row_limit=100,
            ),
            Table("Summary block", [Column("Item"), Column("Count", "int")], [["A", 1]]),
        ],
        notes=["Fixture note."],
    )


def test_pdf_is_paginated_with_header_footer_and_row_note() -> None:
    body = render_pdf(_report(rows=250), compress=False)
    assert body.startswith(b"%PDF")
    pages = len(re.findall(rb"/Type /Page\b", body))
    assert pages >= 3
    text = body.decode("latin-1")
    assert "OPERATIONS-PLATFORM" in text
    assert f"Page 1 of {pages}" in text
    assert f"Page {pages} of {pages}" in text
    assert "Operations-Platform" in text and "Generated Oct 9, 2026 17:30 CDT" in text
    assert "Showing the first 100 of 250 rows" in text
    assert "Department: Production" in text
    assert "Test User \\(fixture\\)" in text or "Test User (fixture)" in text
    # Landscape letter.
    assert "/MediaBox [ 0 0 792 612 ]" in text


def test_pdf_handles_empty_report() -> None:
    empty = _report(rows=0)
    assert render_pdf(empty).startswith(b"%PDF")


def test_excel_types_and_layout() -> None:
    workbook = load_workbook(io.BytesIO(render_excel(_report(rows=5))))
    assert workbook.sheetnames == ["Summary", "Register", "Analysis"]
    summary = workbook["Summary"]
    assert summary["A1"].value == "OPERATIONS-PLATFORM"
    assert summary["A2"].value == "Fixture Report"
    labels = {
        summary.cell(row=r, column=1).value: summary.cell(row=r, column=2) for r in range(1, 40)
    }
    assert labels["Report Name"].value == "Fixture Report"
    assert labels["Module"].value == "Quality · Fixtures"
    assert labels["Period"].value == "Year to Date · Jan 1 – Oct 9, 2026"
    assert isinstance(labels["Generated At"].value, dt.datetime)
    assert labels["Generated At"].value == dt.datetime(2026, 10, 9, 17, 30)
    assert labels["Generated By"].value == TESTER_NAME
    assert labels["Cost"].value == 1234.5 and "$" in labels["Cost"].number_format
    assert labels["Share"].value == 0.125 and labels["Share"].number_format == "0.0%"
    assert labels["Not recorded"].value == "Not recorded"

    sheet = workbook["Register"]
    assert sheet.freeze_panes == "A6"
    assert sheet.auto_filter.ref == "A5:E10"
    assert [c.value for c in sheet[5]] == ["Number", "Date", "Cost", "Share", "Days"]
    first = sheet[6]
    assert isinstance(first[1].value, dt.datetime) and first[1].number_format == "yyyy-mm-dd"
    assert first[2].value is None  # not recorded stays empty, never 0
    second = sheet[7]
    assert second[2].value == 10.25 and "$" in second[2].number_format
    assert second[3].value == 0.5 and second[3].number_format == "0.0%"
    assert second[4].value == 1 and isinstance(second[4].value, int)
    assert sheet.max_row == 10


def test_excel_sheet_names_are_unique_and_valid() -> None:
    report = _report()
    tables = [
        Table("A", [Column("X")], [["1"]], sheet="Data: [2026]/Q1*"),
        Table("B", [Column("X")], [["2"]], sheet="Data: [2026]/Q1*"),
    ]
    workbook = load_workbook(
        io.BytesIO(render_excel(Report(**{**report.__dict__, "tables": tables, "charts": []})))
    )
    assert workbook.sheetnames == ["Summary", "Data   2026  Q1", "Data   2026  Q1 (2)"]


# --- Endpoint guards ------------------------------------------------------------------


def _unused_sessions() -> Any:
    raise AssertionError("no data may be read before the request is authorized")


def _client(principal: Any) -> TestClient:
    app = create_app()
    app.dependency_overrides[get_user_principal] = lambda: principal
    app.dependency_overrides[report_router.report_sessions] = lambda: _unused_sessions
    return TestClient(app)


HEADERS = {"X-Requested-With": "operations-platform"}
URL = "/api/v1/reports/{}/export"


def test_export_requires_sign_in() -> None:
    response = _client(ANONYMOUS).post(
        URL.format("car-dashboard"), json={"format": "pdf"}, headers=HEADERS
    )
    assert response.status_code == 401


@pytest.mark.parametrize("key", sorted(REPORTS))
def test_export_requires_view_and_export_permissions(key: str) -> None:
    definition = REPORTS[key]
    body = {"format": definition.formats[0]}
    # Viewing without the export permission, or exporting without view, is refused.
    for granted in (definition.view_permissions, (definition.export_permission,)):
        response = _client(as_user(*granted)).post(URL.format(key), json=body, headers=HEADERS)
        assert response.status_code == 403, granted
        assert response.json()["detail"]["error"] == "permission_denied"


def test_export_rejects_unknown_report_and_bad_input() -> None:
    every = as_user(*P)
    client = _client(every)
    assert (
        client.post(URL.format("nope"), json={"format": "pdf"}, headers=HEADERS).status_code == 404
    )
    cases = [
        ("car-register", {"format": "pdf"}),
        ("car-register", {"format": "xlsx", "filters": {"unknown": 1}}),
        ("car-register", {"format": "xlsx", "filters": {"department": ["DROP TABLE"]}}),
        ("car-dashboard", {"format": "docx"}),
        (
            "copq",
            {
                "format": "pdf",
                "period": {"mode": "custom", "from": "2026-01-01", "to": "2026-02-01"},
            },
        ),
        ("safety-incidents", {"format": "pdf", "period": {"mode": "ytd"}}),
        (
            "safety-incidents",
            {"format": "pdf", "period": {"mode": "year", "year": 2026}, "filters": {"through": 13}},
        ),
        ("car-dashboard", {"format": "pdf", "period": {"mode": "month", "year": 2026}}),
    ]
    for key, body in cases:
        response = client.post(URL.format(key), json=body, headers=HEADERS)
        assert response.status_code == 422, (key, body)


def test_export_requires_the_application_header() -> None:
    app = create_app()
    app.dependency_overrides[report_router.report_sessions] = lambda: _unused_sessions
    client = TestClient(app)
    client.cookies.set(sessions.cookie_name(get_settings()), "x" * 40)
    response = client.post(URL.format("car-dashboard"), json={"format": "pdf"})
    assert response.status_code == 403
    assert response.json()["detail"]["error"] == "csrf_check_failed"


def test_filenames_are_sanitized() -> None:
    definition = REPORTS["car-dashboard"]
    assert report_router.filename(definition, "2026-08", "pdf") == (
        "Operations-Platform_CAR-Dashboard_2026-08.pdf"
    )
    assert report_router.filename(definition, '2026/08"; x', "xlsx") == (
        "Operations-Platform_CAR-Dashboard_2026-08-x.xlsx"
    )
