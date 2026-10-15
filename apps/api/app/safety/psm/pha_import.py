"""Import a PHA worksheet into Process Safety Management.

    # Read the workbook (opened read-only, never modified) into a mapping for review:
    python -m app.safety.psm.pha_import extract "PSM Dashboard 1.xlsx" mapping.json

    python -m app.safety.psm.pha_import check mapping.json   # validate only
    python -m app.safety.psm.pha_import plan  mapping.json   # compare with stored data
    python -m app.safety.psm.pha_import apply mapping.json   # write (audited, all or nothing)

The same steps are available in the application (``psm.import``): upload the
workbook, review the validation report and the plan, then apply.

Worksheet rules (``2026 PHA - 300,400,600``; no other sheet is read):

- An integer in ``#`` starts a recommendation. Rows below it belong to it until
  the next integer or the footer (``Total``, ``Rejected``, ``Accepted``, ``Open``,
  ``Closed``, ``In-Progress``), which is never imported.
- Each row is one hazard scenario (Event, Cause, Consequence, Priority/Risk).
- A blank Node inherits the node of the row above within the recommendation; a
  blank Area/Equipment inherits the recommendation's.
- The source header "Recomendation" is read as Recommendation.
- Text is kept as written apart from leading and trailing spaces; the raw cell
  values of every row are stored with the scenario (``source_cells``).
- Management columns (Accepted, Task/Justification, Department, Owner, Status,
  Notes) must be blank: imported recommendations are Pending Review with no
  action, department, owner, email, due date or task. Nothing is inferred.
- The extract must match the sheet's expected reconciliation counts, or nothing
  is imported. Imports are idempotent (stable source keys); stored records are
  never overwritten - differences are reported instead.
"""

import argparse
import datetime as dt
import hashlib
import io
import json
import sys
import uuid
from collections import Counter
from collections.abc import Sequence
from dataclasses import dataclass, field
from pathlib import Path
from typing import Annotated, Any, Literal

from pydantic import ConfigDict, Field
from sqlalchemy import select

from app.audit.recorder import AuditChange
from app.core.schemas import CamelModel
from app.quality.cost.records import Actor
from app.safety.models import Area
from app.safety.psm.models import (
    PsmProcessNode,
    PsmRecommendation,
    PsmScenario,
    PsmStudy,
    PsmStudyType,
)
from app.safety.psm.reference import SYSTEM_IMPORT_ACTOR, display_id
from app.safety.psm.service import (
    RECOMMENDATION,
    SCENARIO,
    STUDY,
    name_key,
    recommendation_key,
    scenario_key,
    study_key,
)
from app.safety.psm.store import IMPORT_LOCK_KEY, NUMBER_LOCK_KEY, PsmStore

SHEET = "2026 PHA - 300,400,600"
HEADERS = (
    "#",
    "Area/ Equipment",
    "Recomendation",
    "Consequence",
    "Cause",
    "Event",
    "Node",
    "Priority/ Risk",
    "Accepted (Y/N)",
    "Task/ Justification",
    "Department",
    "Owner",
    "Status",
    "Notes",
)
COLUMNS = "ABCDEFGHIJKLMN"
MANAGEMENT_COLUMNS = "IJKLMN"
FOOTER_LABELS = {"total", "rejected", "accepted", "open", "closed", "in-progress"}
RISK_CODES = {"high": "high", "medium": "medium", "low": "low"}


@dataclass(frozen=True)
class Expected:
    recommendations: int
    scenarios: int
    risks: dict[str, int]
    per_recommendation: dict[int, int]


@dataclass(frozen=True)
class StudyProfile:
    study_code: str
    title: str
    year: int
    study_type: str
    recommendation_prefix: str
    area_codes: tuple[str, ...]
    expected: Expected


# The reconciliation the 2026 sheet must meet (from the reviewed import request).
PROFILES = {
    SHEET: StudyProfile(
        study_code="PHA-2026-300-400-600",
        title=SHEET,
        year=2026,
        study_type="pha",
        recommendation_prefix="PHA-2026",
        area_codes=("300", "400", "600"),
        expected=Expected(
            recommendations=20,
            scenarios=30,
            risks={"high": 0, "medium": 18, "low": 12},
            per_recommendation={
                1: 3,
                2: 2,
                3: 1,
                4: 1,
                5: 1,
                6: 2,
                7: 1,
                8: 1,
                9: 1,
                10: 4,
                11: 1,
                12: 1,
                13: 1,
                14: 1,
                15: 1,
                16: 2,
                17: 2,
                18: 2,
                19: 1,
                20: 1,
            },
        ),
    )
}


class _Model(CamelModel):
    model_config = ConfigDict(extra="forbid", frozen=True)


Text = Annotated[str, Field(min_length=1, max_length=8000)]


class MappedScenario(_Model):
    row: int
    event: Text
    cause: Text
    consequence: Text
    risk: Literal["high", "medium", "low"]
    risk_source: str
    area_equipment: Annotated[str, Field(min_length=1, max_length=200)]
    node: Annotated[str, Field(min_length=1, max_length=1000)]
    cells: dict[str, Any]


class MappedRecommendation(_Model):
    number: Annotated[int, Field(ge=1, le=9999)]
    row_start: int
    row_end: int
    area_equipment: Annotated[str, Field(min_length=1, max_length=200)]
    recommendation_text: Text
    node: Annotated[str, Field(min_length=1, max_length=1000)]
    scenarios: Annotated[list[MappedScenario], Field(min_length=1)]


class MappedSource(_Model):
    file_name: str
    sheet_name: str
    sha256: Annotated[str, Field(pattern=r"^[0-9a-f]{64}$")]


class MappedStudy(_Model):
    study_code: str
    title: str
    year: int
    study_type: str
    recommendation_prefix: str
    area_codes: list[str]


class PhaMapping(_Model):
    source: MappedSource
    study: MappedStudy
    recommendations: list[MappedRecommendation]
    trimmed_values: int = 0
    notes: list[str] = []

    @property
    def study_source_key(self) -> str:
        return f"psm-workbook/{self.source.file_name}/{self.source.sheet_name}"

    def recommendation_source_key(self, number: int) -> str:
        return f"{self.study_source_key}/rec/{number}"


class ImportValidationError(ValueError):
    """The worksheet does not meet the import rules. Nothing was written."""

    def __init__(self, errors: list[str]) -> None:
        super().__init__("; ".join(errors))
        self.errors = errors


# Extraction ----------------------------------------------------------------------------


def _raw(value: Any) -> Any:
    if isinstance(value, dt.datetime | dt.date):
        return value.isoformat()
    return value


def _text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract(
    content: bytes, file_name: str, sheet: str = SHEET, profile: StudyProfile | None = None
) -> PhaMapping:
    """Read the worksheet into a mapping. Raises ImportValidationError, listing
    every problem found, when the sheet breaks a rule."""
    from openpyxl import load_workbook

    profile = profile or PROFILES.get(sheet)
    if profile is None:
        raise ImportValidationError([f"No import profile for worksheet “{sheet}”."])
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        if sheet not in workbook.sheetnames:
            raise ImportValidationError([f"The workbook has no worksheet named “{sheet}”."])
        rows = [
            [cell for cell in row[: len(COLUMNS)]]
            for row in workbook[sheet].iter_rows(min_row=1, max_col=len(COLUMNS), values_only=True)
        ]
    finally:
        workbook.close()
    errors: list[str] = []
    header = [(_text(v) or "") for v in (rows[0] if rows else [])]
    header += [""] * (len(COLUMNS) - len(header))
    for column, expected, found in zip(COLUMNS, HEADERS, header, strict=True):
        if found.lower() != expected.lower():
            errors.append(f"Header {column}1 reads “{found}”, expected “{expected}”.")
    if errors:
        raise ImportValidationError(errors)

    trimmed = 0
    recommendations: list[dict[str, Any]] = []
    current: dict[str, Any] | None = None
    footer_total: Any = None
    for index, values in enumerate(rows[1:], start=2):
        values = list(values) + [None] * (len(COLUMNS) - len(values))
        cells = {
            f"{c}{index}": _raw(v) for c, v in zip(COLUMNS, values, strict=True) if v is not None
        }
        label = (_text(values[1]) or "").lower()
        if _text(values[0]) is None and label in FOOTER_LABELS:
            if label == "total":
                footer_total = values[2]
            current = None
            break
        if not cells:
            continue
        for value in values:
            if isinstance(value, str) and value != value.strip():
                trimmed += 1
        number = values[0]
        if number is not None:
            if isinstance(number, float) and number.is_integer():
                number = int(number)
            if not isinstance(number, int) or isinstance(number, bool):
                errors.append(f"Row {index}: “#” must be a whole number, found “{number}”.")
                current = None
                continue
            text = _text(values[2])
            area = _text(values[1])
            if text is None:
                errors.append(f"Row {index}: recommendation {number} has no recommendation text.")
            if area is None:
                errors.append(f"Row {index}: recommendation {number} has no Area/Equipment.")
            current = {
                "number": number,
                "row_start": index,
                "row_end": index,
                "area_equipment": area or "",
                "recommendation_text": text or "",
                "node": _text(values[6]) or "",
                "scenarios": [],
            }
            recommendations.append(current)
        elif current is None:
            errors.append(f"Row {index}: data outside any numbered recommendation.")
            continue
        elif _text(values[2]) is not None:
            errors.append(
                f"Row {index}: recommendation text without a number in “#” "
                f"(inside recommendation {current['number']})."
            )
        current["row_end"] = index
        for column in MANAGEMENT_COLUMNS:
            value = values[COLUMNS.index(column)]
            if _text(value) is not None:
                errors.append(
                    f"Cell {column}{index} has a value (“{value}”); management columns must be "
                    "blank for this import."
                )
        previous = current["scenarios"][-1] if current["scenarios"] else None
        node = _text(values[6]) or (previous["node"] if previous else current["node"])
        area = _text(values[1]) or current["area_equipment"]
        risk_source = values[7]
        risk = RISK_CODES.get((_text(risk_source) or "").lower())
        scenario = {
            "row": index,
            "event": _text(values[5]),
            "cause": _text(values[4]),
            "consequence": _text(values[3]),
            "risk": risk,
            "risk_source": "" if risk_source is None else str(risk_source),
            "area_equipment": area,
            "node": node,
            "cells": cells,
        }
        for name, label_text in (
            ("event", "Event"),
            ("cause", "Cause"),
            ("consequence", "Consequence"),
        ):
            if scenario[name] is None:
                errors.append(f"Row {index}: {label_text} is blank.")
        if risk is None:
            errors.append(f"Row {index}: Priority/Risk “{risk_source}” is not High, Medium or Low.")
        if not node:
            errors.append(f"Row {index}: no Node (none above it within the recommendation).")
        current["scenarios"].append(scenario)

    numbers = [r["number"] for r in recommendations]
    if numbers != list(range(1, len(numbers) + 1)):
        errors.append(f"Recommendation numbers are not 1 to {len(numbers)} in order: {numbers}.")
    if footer_total is not None and footer_total != len(recommendations):
        errors.append(
            f"The sheet's Total reads {footer_total}; "
            f"{len(recommendations)} recommendations were found."
        )
    errors += _reconcile(profile.expected, recommendations)
    if errors:
        raise ImportValidationError(errors)
    return PhaMapping.model_validate(
        {
            "source": {
                "file_name": Path(file_name).name,
                "sheet_name": sheet,
                "sha256": hashlib.sha256(content).hexdigest(),
            },
            "study": {
                "study_code": profile.study_code,
                "title": profile.title,
                "year": profile.year,
                "study_type": profile.study_type,
                "recommendation_prefix": profile.recommendation_prefix,
                "area_codes": list(profile.area_codes),
            },
            "recommendations": recommendations,
            "trimmed_values": trimmed,
            "notes": [
                "Generated by app.safety.psm.pha_import from the workbook; review before applying.",
                "Leading and trailing spaces were trimmed from working values; "
                "each scenario keeps its row's raw cells.",
            ],
        }
    )


def _reconcile(expected: Expected, recommendations: list[dict[str, Any]]) -> list[str]:
    errors = []
    scenarios = [s for r in recommendations for s in r["scenarios"]]
    if len(recommendations) != expected.recommendations:
        errors.append(
            f"Found {len(recommendations)} recommendations; expected {expected.recommendations}."
        )
    if len(scenarios) != expected.scenarios:
        errors.append(f"Found {len(scenarios)} hazard scenarios; expected {expected.scenarios}.")
    risks = Counter(s["risk"] for s in scenarios)
    for risk, count in expected.risks.items():
        if risks[risk] != count:
            errors.append(f"Found {risks[risk]} {risk.title()} scenarios; expected {count}.")
    for r in recommendations:
        want = expected.per_recommendation.get(r["number"])
        if want is not None and len(r["scenarios"]) != want:
            errors.append(
                f"Recommendation {r['number']} has {len(r['scenarios'])} scenarios; "
                f"expected {want}."
            )
    return errors


def validate(mapping: PhaMapping) -> list[str]:
    """Re-check a mapping file against the profile (it may have been edited)."""
    profile = PROFILES.get(mapping.source.sheet_name)
    if profile is None:
        return [f"No import profile for worksheet “{mapping.source.sheet_name}”."]
    errors = _reconcile(profile.expected, [r.model_dump() for r in mapping.recommendations])
    numbers = [r.number for r in mapping.recommendations]
    if numbers != list(range(1, len(numbers) + 1)):
        errors.append("Recommendation numbers must run from 1 in order.")
    if mapping.study.recommendation_prefix != profile.recommendation_prefix:
        errors.append("The recommendation prefix differs from the worksheet's profile.")
    return errors


# Plan and apply ---------------------------------------------------------------------------


@dataclass
class Plan:
    study_exists: bool = False
    new_recommendations: list[int] = field(default_factory=list)
    unchanged_recommendations: list[int] = field(default_factory=list)
    new_scenarios: int = 0
    new_nodes: list[str] = field(default_factory=list)
    differences: list[str] = field(default_factory=list)
    conflicts: list[str] = field(default_factory=list)

    def summary(self) -> dict[str, Any]:
        return {
            "studyExists": self.study_exists,
            "newRecommendations": len(self.new_recommendations),
            "unchangedRecommendations": len(self.unchanged_recommendations),
            "newScenarios": self.new_scenarios,
            "newProcessNodes": self.new_nodes,
            "differences": self.differences,
            "conflicts": self.conflicts,
        }


SCENARIO_COMPARED = ("event", "cause", "consequence", "risk", "area_equipment")


def plan(store: PsmStore, mapping: PhaMapping) -> Plan:
    result = Plan()
    session = store.session
    study = session.scalar(select(PsmStudy).where(PsmStudy.source_key == mapping.study_source_key))
    nodes = {n.name_key: n for n in store.process_nodes()}
    wanted_nodes = {r.node for r in mapping.recommendations} | {
        s.node for r in mapping.recommendations for s in r.scenarios
    }
    result.new_nodes = sorted(n for n in wanted_nodes if name_key(n) not in nodes)
    for node in wanted_nodes:
        existing = nodes.get(name_key(node))
        if existing is not None and existing.name != node:
            result.differences.append(
                f"Process node “{node}” matches the stored node “{existing.name}”; "
                "the stored name is used."
            )
    if study is None:
        clash = session.scalar(
            select(PsmStudy).where(
                (PsmStudy.study_code == mapping.study.study_code)
                | (PsmStudy.recommendation_prefix == mapping.study.recommendation_prefix)
            )
        )
        if clash is not None:
            result.conflicts.append(
                f"Study {clash.study_code} already uses this study code or recommendation prefix "
                "but was not imported from this worksheet."
            )
        result.new_recommendations = [r.number for r in mapping.recommendations]
        result.new_scenarios = sum(len(r.scenarios) for r in mapping.recommendations)
        return result
    result.study_exists = True
    if study.title != mapping.study.title or study.year != mapping.study.year:
        result.differences.append(
            "The stored study's title or year differs from the worksheet; kept as stored."
        )
    stored = {
        r.recommendation_number: r
        for r in session.scalars(
            select(PsmRecommendation).where(PsmRecommendation.study_id == study.id)
        )
    }
    for mapped in mapping.recommendations:
        rec = stored.get(mapped.number)
        if rec is None:
            result.new_recommendations.append(mapped.number)
            result.new_scenarios += len(mapped.scenarios)
            continue
        if rec.source_key != mapping.recommendation_source_key(mapped.number):
            result.conflicts.append(
                f"{rec.display_id} exists but was not imported from row {mapped.row_start}."
            )
            continue
        result.unchanged_recommendations.append(mapped.number)
        label = rec.display_id
        if rec.area_equipment != mapped.area_equipment:
            result.differences.append(f"{label}: Area/Equipment differs (stored value kept).")
        if rec.recommendation_text != mapped.recommendation_text:
            result.differences.append(f"{label}: recommendation text differs (stored value kept).")
        scenarios = {s.source_row: s for s in store.scenarios_of(rec.id) if s.source_row}
        for mapped_scenario in mapped.scenarios:
            scenario = scenarios.get(mapped_scenario.row)
            if scenario is None:
                result.differences.append(
                    f"{label}: the scenario from row {mapped_scenario.row} is not stored "
                    "(not added)."
                )
                continue
            changed = [
                name
                for name in SCENARIO_COMPARED
                if getattr(scenario, name) != getattr(mapped_scenario, name)
            ]
            if changed:
                result.differences.append(
                    f"{label}: scenario from row {mapped_scenario.row} differs in "
                    f"{', '.join(changed)} (stored values kept)."
                )
    return result


@dataclass(frozen=True)
class ApplyResult:
    study_id: int
    recommendations: int
    scenarios: int
    nodes: int
    plan: Plan


def apply(store: PsmStore, mapping: PhaMapping, actor: Actor) -> ApplyResult:
    """Create what is missing, all in one transaction. Stored records are never
    changed. Refused (nothing written) on validation errors or conflicts."""
    errors = validate(mapping)
    if errors:
        raise ImportValidationError(errors)
    change_set = uuid.uuid4()
    try:
        store.lock(IMPORT_LOCK_KEY)
        store.lock(NUMBER_LOCK_KEY)
        result_plan = plan(store, mapping)
        if result_plan.conflicts:
            raise ImportValidationError(result_plan.conflicts)
        session = store.session
        changes: list[AuditChange] = []
        stamps = {
            "created_at": actor.now,
            "created_by": actor.actor_id,
            "updated_at": actor.now,
            "updated_by": actor.actor_id,
        }
        nodes = {n.name_key: n for n in store.process_nodes()}
        order = max((n.display_order for n in nodes.values()), default=0)
        created_nodes = 0
        in_sheet_order = dict.fromkeys(
            name
            for r in mapping.recommendations
            for name in (r.node, *(s.node for s in r.scenarios))
        )
        for name in in_sheet_order:
            key = name_key(name)
            if key in nodes:
                continue
            order += 1
            node = PsmProcessNode(name=name, name_key=key, display_order=order, **stamps)
            session.add(node)
            session.flush()
            nodes[key] = node
            created_nodes += 1
            changes.append(
                AuditChange(
                    "create",
                    "safety.psm_reference",
                    f"psm/reference/psm_process_nodes/{node.id}",
                    None,
                    {"name": name, "event": "imported", "source": mapping.study_source_key},
                )
            )
        study = session.scalar(
            select(PsmStudy).where(PsmStudy.source_key == mapping.study_source_key)
        )
        if study is None:
            study_type = session.scalar(
                select(PsmStudyType).where(PsmStudyType.code == mapping.study.study_type)
            )
            areas = list(
                session.scalars(select(Area).where(Area.code.in_(mapping.study.area_codes)))
            )
            if study_type is None or len(areas) != len(mapping.study.area_codes):
                raise ImportValidationError(["The study type or plant areas are not set up."])
            study = PsmStudy(
                study_code=mapping.study.study_code,
                title=mapping.study.title,
                year=mapping.study.year,
                study_type_id=study_type.id,
                recommendation_prefix=mapping.study.recommendation_prefix,
                source_file_name=mapping.source.file_name,
                source_sheet_name=mapping.source.sheet_name,
                source_reference=f"{mapping.source.file_name} › {mapping.source.sheet_name} "
                f"(SHA-256 {mapping.source.sha256})",
                source_key=mapping.study_source_key,
                **stamps,
            )
            session.add(study)
            session.flush()
            store.set_study_areas(study.id, [a.id for a in areas])
            changes.append(
                AuditChange(
                    "create",
                    STUDY,
                    study_key(study.id),
                    None,
                    {
                        "study_code": study.study_code,
                        "title": study.title,
                        "source": study.source_reference,
                        "event": "imported",
                    },
                )
            )
        created_scenarios = 0
        wanted = set(result_plan.new_recommendations)
        for mapped in mapping.recommendations:
            if mapped.number not in wanted:
                continue
            rec = PsmRecommendation(
                study_id=study.id,
                recommendation_number=mapped.number,
                display_id=display_id(study.recommendation_prefix, mapped.number),
                area_equipment=mapped.area_equipment,
                recommendation_text=mapped.recommendation_text,
                process_node_id=nodes[name_key(mapped.node)].id,
                source_row_start=mapped.row_start,
                source_row_end=mapped.row_end,
                source_key=mapping.recommendation_source_key(mapped.number),
                **stamps,
            )
            session.add(rec)
            session.flush()
            changes.append(
                AuditChange(
                    "create",
                    RECOMMENDATION,
                    recommendation_key(rec.id),
                    None,
                    {
                        "display_id": rec.display_id,
                        "area_equipment": rec.area_equipment,
                        "recommendation_text": rec.recommendation_text,
                        "disposition": "pending_review",
                        "source_rows": f"{mapped.row_start}-{mapped.row_end}",
                        "event": "imported",
                    },
                )
            )
            for number, mapped_scenario in enumerate(mapped.scenarios, 1):
                scenario = PsmScenario(
                    recommendation_id=rec.id,
                    sequence_number=number,
                    event=mapped_scenario.event,
                    cause=mapped_scenario.cause,
                    consequence=mapped_scenario.consequence,
                    risk=mapped_scenario.risk,
                    risk_source=mapped_scenario.risk_source,
                    area_equipment=mapped_scenario.area_equipment,
                    process_node_id=nodes[name_key(mapped_scenario.node)].id,
                    source_row=mapped_scenario.row,
                    source_cells=mapped_scenario.cells,
                    **stamps,
                )
                session.add(scenario)
                session.flush()
                created_scenarios += 1
                changes.append(
                    AuditChange(
                        "create",
                        SCENARIO,
                        scenario_key(rec.id, scenario.id),
                        None,
                        {
                            "source_row": mapped_scenario.row,
                            "risk": scenario.risk,
                            "event": "imported",
                        },
                    )
                )
        store.audit(
            actor_id=actor.actor_id, change_set_id=change_set, at=actor.now, changes=changes
        )
        store.commit()
        return ApplyResult(study.id, len(wanted), created_scenarios, created_nodes, result_plan)
    except Exception:
        store.rollback()
        raise


# Command line ------------------------------------------------------------------------------


def _load(path: Path) -> PhaMapping:
    return PhaMapping.model_validate_json(path.read_text(encoding="utf-8"))


def _print_plan(result: Plan) -> None:
    summary = result.summary()
    print(f"Study already imported: {'yes' if summary['studyExists'] else 'no'}")
    print(f"Recommendations to add: {summary['newRecommendations']}")
    print(f"Recommendations already stored (left unchanged): {summary['unchangedRecommendations']}")
    print(f"Hazard scenarios to add: {summary['newScenarios']}")
    print(f"Process nodes to add: {len(summary['newProcessNodes'])}")
    for line in summary["differences"]:
        print(f"  difference: {line}")
    for line in summary["conflicts"]:
        print(f"  CONFLICT: {line}")


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(prog="python -m app.safety.psm.pha_import")
    commands = parser.add_subparsers(dest="command", required=True)
    extract_parser = commands.add_parser("extract")
    extract_parser.add_argument("workbook", type=Path)
    extract_parser.add_argument("output", type=Path)
    for name in ("check", "plan", "apply"):
        commands.add_parser(name).add_argument("mapping", type=Path)
    args = parser.parse_args(argv)
    if args.command == "extract":
        try:
            mapping = extract(args.workbook.read_bytes(), args.workbook.name)
        except ImportValidationError as error:
            print("Validation failed; nothing written:", file=sys.stderr)
            for line in error.errors:
                print(f"  - {line}", file=sys.stderr)
            return 1
        args.output.write_text(
            json.dumps(mapping.model_dump(mode="json", by_alias=True), indent=2, ensure_ascii=False)
            + "\n",
            encoding="utf-8",
        )
        count = sum(len(r.scenarios) for r in mapping.recommendations)
        print(
            f"Wrote {len(mapping.recommendations)} recommendations, {count} scenarios "
            f"to {args.output}."
        )
        return 0
    mapping = _load(args.mapping)
    errors = validate(mapping)
    if errors:
        print("Validation failed; nothing written:", file=sys.stderr)
        for line in errors:
            print(f"  - {line}", file=sys.stderr)
        return 1
    if args.command == "check":
        print("The mapping meets the import rules.")
        return 0
    from app.db.session import get_sessionmaker

    with get_sessionmaker()() as session:
        store = PsmStore(session)
        if args.command == "plan":
            _print_plan(plan(store, mapping))
            return 0
        try:
            result = apply(store, mapping, Actor(SYSTEM_IMPORT_ACTOR, dt.datetime.now(dt.UTC)))
        except ImportValidationError as error:
            print("Refused; nothing written:", file=sys.stderr)
            for line in error.errors:
                print(f"  - {line}", file=sys.stderr)
            return 1
    _print_plan(result.plan)
    print(
        f"Imported {result.recommendations} recommendations, {result.scenarios} scenarios and "
        f"{result.nodes} process nodes."
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
